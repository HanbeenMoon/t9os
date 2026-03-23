"""T9 OS 파일 파서 유틸리티 — t9_seed.py에서 분리"""
import re, json
from pathlib import Path

def parse_file(filepath):
    """Parse any supported file. Returns (meta_dict, body_text)."""
    fp = Path(filepath)
    ext = fp.suffix.lower()
    if ext == ".md":
        return _parse_md_content(fp.read_text(encoding="utf-8", errors="replace"))
    elif ext in (".txt", ".csv", ".log"):
        return {}, fp.read_text(encoding="utf-8", errors="replace")
    elif ext == ".docx":
        try:
            from docx import Document
            doc = Document(str(fp))
            return {}, "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception:
            return {}, f"[DOCX] {fp.name}"
    elif ext == ".pdf":
        try:
            import fitz
            doc = fitz.open(str(fp))
            text = "\n".join(page.get_text() for page in doc[:10])
            doc.close()
            return {}, text
        except Exception:
            return {}, f"[PDF] {fp.name}"
    elif ext == ".xlsx":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(fp), read_only=True)
            lines = []
            for ws in wb.worksheets[:3]:
                for row in ws.iter_rows(max_row=20, values_only=True):
                    lines.append(" ".join(str(c) for c in row if c is not None))
            wb.close()
            return {}, "\n".join(lines)
        except Exception:
            return {}, f"[XLSX] {fp.name}"
    elif ext in (".jpg", ".jpeg", ".png", ".svg", ".mp4", ".zip", ".hwp", ".exe"):
        return {}, f"[{ext[1:].upper()}] {fp.name}"
    else:
        return {}, f"[{ext[1:].upper() if ext else 'UNKNOWN'}] {fp.name}"

def _parse_md_content(text):
    meta, body = {}, text
    m = re.match(r'^---\s*\n(.*?)\n---\s*\n(.*)', text, re.DOTALL)
    if m:
        body = m.group(2)
        for line in m.group(1).split("\n"):
            line = line.strip()
            if ":" in line:
                key, _, val = line.partition(":")
                key, val = key.strip(), val.strip()
                if val.startswith("[") and val.endswith("]"):
                    val = [v.strip().strip("'\"") for v in val[1:-1].split(",") if v.strip()]
                elif val.lower() in ("true", "false"):
                    val = val.lower() == "true"
                elif val.replace(".", "").replace("-", "").isdigit():
                    try: val = float(val) if "." in val else int(val)
                    except: pass
                elif val.startswith(("'", '"')) and val.endswith(("'", '"')):
                    val = val[1:-1]
                meta[key] = val
    return meta, body.strip()

def parse_md(filepath):
    """Legacy wrapper — delegates to _parse_md_content."""
    return _parse_md_content(Path(filepath).read_text(encoding="utf-8", errors="replace"))

def write_md(filepath, meta, body):
    lines = ["---"]
    for k, v in meta.items():
        if isinstance(v, list): lines.append(f"{k}: [{', '.join(str(x) for x in v)}]")
        elif isinstance(v, bool): lines.append(f"{k}: {'true' if v else 'false'}")
        elif isinstance(v, dict): lines.append(f"{k}: {json.dumps(v, ensure_ascii=False)}")
        else: lines.append(f"{k}: {v}")
    lines += ["---", "", body]
    Path(filepath).write_bytes("\n".join(lines).encode("utf-8"))
